from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from edu_ops.weekly_report.production import (
    _excel_binding_check,
    discover_sources,
    main,
    resolve_periods,
    run_production,
    week_delta,
)


ASSET_ROOT = Path("/Users/macos/Documents/03-周报数据/数学组周数据统计")


def _real_source_manifest():
    return discover_sources(ASSET_ROOT)


def test_source_manifest_and_period_gate_find_latest_complete_real_period():
    manifest = _real_source_manifest()
    assert manifest["files"]
    roles = {item["role"] for item in manifest["files"]}
    assert {"WPS", "TMS", "FINAL_GOLDEN"}.issubset(roles)
    resolution = resolve_periods(manifest)
    assert resolution["latest_available_period"] == (2026, 6, 4)
    assert resolution["latest_complete_period"] == (2026, 6, 4)
    assert resolution["completeness"] == "COMPLETE"


def test_next_week_fixture_is_discovered_without_rule_changes(tmp_path: Path):
    source = tmp_path / "inbox"
    source.mkdir()
    for name in ("二校数学组数据汇总-六月第四周.xls", "数学组数据统计表-宣城二校6月第4周.xls", "排课列表_06月01日到06月28日_202606301011.xls"):
        shutil.copy2(ASSET_ROOT / name, source / name)
    initial = resolve_periods(discover_sources(source))
    assert initial["latest_complete_period"] == (2026, 6, 4)
    # A new period is represented by the same real schema, not a synthetic row.
    shutil.copy2(source / "二校数学组数据汇总-六月第四周.xls", source / "二校数学组数据汇总-七月第一周.xls")
    shutil.copy2(source / "数学组数据统计表-宣城二校6月第4周.xls", source / "数学组数据统计表-宣城二校7月第1周.xls")
    shutil.copy2(source / "排课列表_06月01日到06月28日_202606301011.xls", source / "排课列表_07月01日到07月07日_202607071011.xls")
    updated = resolve_periods(discover_sources(source))
    assert updated["latest_available_period"] == (2026, 7, 1)
    assert updated["latest_complete_period"] == (2026, 7, 1)


def test_partial_latest_period_does_not_fall_back_to_an_older_final(tmp_path: Path):
    source = tmp_path / "inbox"
    source.mkdir()
    for name in ("二校数学组数据汇总-六月第四周.xls", "数学组数据统计表-宣城二校6月第4周.xls", "排课列表_06月01日到06月28日_202606301011.xls"):
        shutil.copy2(ASSET_ROOT / name, source / name)
    shutil.copy2(source / "二校数学组数据汇总-六月第四周.xls", source / "二校数学组数据汇总-七月第一周.xls")

    resolution = resolve_periods(discover_sources(source))
    assert resolution["latest_available_period"] == (2026, 7, 1)
    assert resolution["latest_complete_period"] == (2026, 6, 4)
    assert resolution["completeness"] == "PARTIAL"

    output = tmp_path / "out"
    prior = output / "weekly_report_2026-06-w4.xlsx"
    output.mkdir()
    prior.write_bytes(b"previous-success")
    result = run_production(source, output)
    assert result["status"] == "PARTIAL"
    assert not (output / "weekly_report_snapshot.json").exists()
    assert prior.exists()


def test_real_template_mapping_and_tamper_gate(tmp_path: Path):
    output = tmp_path / "out"
    result = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert result["status"] == "PASS"
    snapshot = json.loads((output / "weekly_report_snapshot.json").read_text(encoding="utf-8"))
    dated = output / "weekly_report_2026-06-w4.xlsx"
    assert _excel_binding_check(dated, snapshot)["status"] == "PASS"
    book = load_workbook(dated)
    book["学生"]["C7"] = 999
    book.save(dated)
    assert _excel_binding_check(dated, snapshot)["status"] == "FAILED"


def test_main_failure_is_structured_without_secondary_traceback(tmp_path: Path, monkeypatch, capsys):
    import edu_ops.weekly_report.production as production

    def explode(*args, **kwargs):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(production, "_run_production", explode)
    output = tmp_path / "out"
    assert main(["--source-root", str(tmp_path / "missing"), "--output-root", str(output)]) == 2
    captured = capsys.readouterr()
    assert json.loads(captured.out)["status"] == "FAILED"
    assert "Traceback" not in captured.err
    assert json.loads((output / "PIPELINE_STATUS.json").read_text())["status"] == "FAILED"


def test_manifest_hash_is_stable_and_bound_to_snapshot(tmp_path: Path):
    first = discover_sources(ASSET_ROOT)
    second = discover_sources(ASSET_ROOT)
    assert first["manifest_hash"] == second["manifest_hash"]
    output = tmp_path / "out"
    result = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert result["snapshot"]["manifest_hash"] == first["manifest_hash"]


def test_production_dry_run_is_idempotent_and_exports_reopenable_excel(tmp_path: Path):
    output = tmp_path / "out"
    first = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    first_snapshot = json.loads((output / "weekly_report_snapshot.json").read_text(encoding="utf-8"))
    second = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    second_snapshot = json.loads((output / "weekly_report_snapshot.json").read_text(encoding="utf-8"))
    assert first["status"] == second["status"] == "PASS"
    assert first_snapshot["business_fingerprint"] == second_snapshot["business_fingerprint"]
    assert second_snapshot["idempotency"]["same_input_business_fingerprint"] is True
    workbook = load_workbook(output / "weekly_report_2026-06-w4.xlsx", data_only=False)
    assert {"学生", "组课时生产", "教师", "生产元数据"}.issubset(workbook.sheetnames)
    assert all("#REF!" not in str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)


def test_empty_output_repeated_identical_runs_are_byte_identical(tmp_path: Path):
    output = tmp_path / "out"
    first = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert first["status"] == "PASS"
    snapshot_path = output / "weekly_report_snapshot.json"
    first_bytes = snapshot_path.read_bytes()
    first_snapshot = json.loads(first_bytes)
    assert first_snapshot["week_over_week"]["status"] == "NO_PREVIOUS_PERIOD"

    second = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert second["status"] == "PASS"
    second_bytes = snapshot_path.read_bytes()
    assert second_bytes == first_bytes
    assert json.loads(second_bytes)["week_over_week"] == first_snapshot["week_over_week"]


def test_excel_semantic_regression_mapping_is_not_just_expected_cell_self_check(tmp_path: Path):
    output = tmp_path / "out"
    result = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert result["status"] == "PASS"
    snapshot = json.loads((output / "weekly_report_snapshot.json").read_text(encoding="utf-8"))
    workbook = load_workbook(output / "weekly_report_2026-06-w4.xlsx", data_only=False)
    production = workbook["组课时生产"]
    teachers = workbook["教师"]

    # Week 4 is row 7: H/L are student counts, I/N are production KS,
    # M is class count, and Q is the formal teacher count.
    assert production["H7"].value == 94  # 1v1 初小 + 高中; 双三 is overlap only
    assert production["I7"].value == snapshot["production"]["tms"]["one_to_one_ks"] == 219
    assert production["L7"].value == 169  # 班课初小 + 高中
    assert production["M7"].value == snapshot["students"]["class"]["classes"] == 39
    assert production["N7"].value == snapshot["production"]["tms"]["class_ks"] == 411
    assert production["Q7"].value == snapshot["production"]["golden"]["teacher_count"] == 18

    # Teacher student counts include both parent grade buckets; high-school
    # rows must not silently disappear from the teacher sheet.
    assert teachers["C5"].value == 4  # 胡长春: 0 初小 + 4 高中
    assert teachers["F5"].value == 11  # 胡长春: 0 初小 + 11 高中

    # The source snapshot has no month-to-date student/KS breakdown.  The
    # month-total row therefore keeps only its explicitly sourced B value.
    assert production["B9"].value == snapshot["production"]["golden"]["month_hours"]
    for cell in ("C9", "H9", "I9", "L9", "M9", "N9", "Q9"):
        assert production[cell].value is None
    for row in (7, 11):
        for column in ("D", "E", "F", "G", "J", "K", "O", "P"):
            assert production[f"{column}{row}"].value is None


def test_week_delta_is_explicit_and_bounded():
    current = {"period": {"year": 2026, "month": 6, "week": 4}, "students": {"single_subject_total": 12}, "production": {"tms": {"one_to_one_ks": 10, "class_ks": 20, "week_hours_equivalent": 16}}, "teachers": [{"name": "A"}, {"name": "B"}]}
    previous = {"period": {"year": 2026, "month": 6, "week": 3}, "students": {"single_subject_total": 10}, "production": {"tms": {"one_to_one_ks": 8, "class_ks": 20, "week_hours_equivalent": 14}}, "teachers": [{"name": "A"}]}
    result = week_delta(current, previous)
    assert result["status"] == "COMPUTED"
    assert result["metrics"]["students.single_subject_total"]["delta"] == 2
    assert result["metrics"]["teachers.count"]["delta"] == 1


def test_same_period_replacement_does_not_self_compare_as_zero(tmp_path: Path):
    output = tmp_path / "out"
    first = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert first["status"] == "PASS"
    snapshot_path = output / "weekly_report_snapshot.json"
    previous = json.loads(snapshot_path.read_text(encoding="utf-8"))
    previous["students"]["single_subject_total"] += 1
    previous["business_fingerprint"] = "different-same-period-baseline"
    previous["week_over_week"] = {"status": "COMPUTED", "metrics": {"students.single_subject_total": {"delta": 0}}}
    snapshot_path.write_text(json.dumps(previous, ensure_ascii=False, indent=2), encoding="utf-8")

    second = run_production(ASSET_ROOT, output, template=ASSET_ROOT / "数学组数据统计表基础模板.xlsx")
    assert second["status"] == "PASS"
    refreshed = json.loads(snapshot_path.read_text(encoding="utf-8"))
    assert refreshed["week_over_week"]["status"] == "SAME_PERIOD_BASELINE_UNAVAILABLE"
    assert refreshed["week_over_week"]["metrics"] == {}


def test_cross_period_week_delta_is_idempotent(tmp_path: Path):
    source = tmp_path / "inbox"
    shutil.copytree(ASSET_ROOT, source)
    output = tmp_path / "out"
    first = run_production(source, output)
    assert first["status"] == "PASS"

    shutil.copy2(source / "二校数学组数据汇总-六月第四周.xls", source / "二校数学组数据汇总-七月第一周.xls")
    shutil.copy2(source / "数学组数据统计表-宣城二校6月第4周.xls", source / "数学组数据统计表-宣城二校7月第一周.xls")
    shutil.copy2(source / "排课列表_06月01日到06月28日_202606301011.xls", source / "排课列表_07月01日到07月07日_202607071011.xls")

    second = run_production(source, output)
    assert second["status"] == "PASS"
    snapshot_path = output / "weekly_report_snapshot.json"
    first_new_bytes = snapshot_path.read_bytes()
    first_new = json.loads(first_new_bytes)
    assert first_new["week_over_week"]["status"] == "COMPUTED"
    assert tuple(first_new["week_over_week"]["current_period"]) == (2026, 7, 1)
    assert tuple(first_new["week_over_week"]["previous_period"]) == (2026, 6, 4)

    third = run_production(source, output)
    assert third["status"] == "PASS"
    second_new_bytes = snapshot_path.read_bytes()
    assert hashlib.sha256(second_new_bytes).hexdigest() == hashlib.sha256(first_new_bytes).hexdigest()
    assert json.loads(second_new_bytes)["week_over_week"] == first_new["week_over_week"]


def test_same_period_week_delta_is_not_computed():
    current = {"period": {"year": 2026, "month": 6, "week": 4}, "students": {"single_subject_total": 12}, "teachers": []}
    previous = {"period": {"year": 2026, "month": 6, "week": 4}, "students": {"single_subject_total": 10}, "teachers": []}
    result = week_delta(current, previous)
    assert result["status"] == "SAME_PERIOD_BASELINE_UNAVAILABLE"
    assert result["metrics"] == {}
